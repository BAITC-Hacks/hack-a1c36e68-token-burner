"""File -> Document with numbered clauses.

Clause numbers look like `1.5.`, `2.4.16.`, `10.` (section heading) and sub-items `а.`/`а)`.
A number at the start of a line always opens a clause; an inline number (`... 3.10.Работники`)
only when followed by a capital letter. Every candidate must also be a plausible successor of
the previous clause number, which filters out cross-references, dates and the table of contents.
"""
import re
import sys
from pathlib import Path

from backend.agent.schemas import Clause, Document, Side

NUM_RE = re.compile(r"(?<![\w.])(\d{1,2}(?:\.\d{1,2}){0,3})\.(?!\d)[ \t]*")
SUB_RE = re.compile(r"(?m)^[ \t]*([а-яё])[.)][ \t]+")
TOC_RE = re.compile(r"(?mi)^[ \t]*(оглавление|содержание)[ \t]*$")
CAPITAL = re.compile(r"[А-ЯЁA-Z«]")
TOC_LINE = re.compile(r"(\.{4,}|…{2,}|\s\d{1,3})\s*$")
FOOTER = re.compile(r"\n\s*\d{1,3}\s*$")
LETTERS = "абвгдежзиклмнопрстуфхцчшщэюя"
MAX_GAP = 3  # tolerated jump in numbering, e.g. 5.5.3 -> 5.5.6


# --- text extraction ---


def pdf_pages(path: Path) -> list[str]:
    try:
        import pymupdf

        with pymupdf.open(path) as doc:
            return [page.get_text() for page in doc]
    except Exception:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return [page.extract_text() or "" for page in pdf.pages]


def docx_text(path: Path) -> str:
    """Paragraphs and table cells in body order, with auto-numbering rendered as text."""
    import docx
    from docx.oxml.ns import qn

    document = docx.Document(str(path))
    numbering = _Numbering(document)
    lines = []
    for el in document.element.body.iterchildren():
        if el.tag == qn("w:p"):
            lines.append(numbering.label(el) + _p_text(el))
        elif el.tag == qn("w:tbl"):
            for p in el.iter(qn("w:p")):
                lines.append(_p_text(p))
    return "\n".join(line for line in lines if line.strip())


def _p_text(p) -> str:
    from docx.oxml.ns import qn

    return "".join(t.text or "" for t in p.iter(qn("w:t")))


class _Numbering:
    """Minimal renderer of Word list numbering (decimal and Russian letters)."""

    def __init__(self, document):
        from docx.oxml.ns import qn

        self.qn = qn
        self.levels: dict[str, dict[int, tuple[str, str, int]]] = {}  # numId -> ilvl -> (fmt, text, start)
        self.abstract_of: dict[str, str] = {}
        self.counters: dict[str, list[int]] = {}
        try:
            root = document.part.numbering_part.element
        except Exception:
            return
        abstract = {}
        for an in root.iter(qn("w:abstractNum")):
            lvls = {}
            for lvl in an.iter(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                text = lvl.find(qn("w:lvlText"))
                start = lvl.find(qn("w:start"))
                lvls[int(lvl.get(qn("w:ilvl")))] = (
                    fmt.get(qn("w:val")) if fmt is not None else "decimal",
                    text.get(qn("w:val")) if text is not None else "",
                    int(start.get(qn("w:val"))) if start is not None else 1,
                )
            abstract[an.get(qn("w:abstractNumId"))] = lvls
        for num in root.iter(qn("w:num")):
            aid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            self.levels[num.get(qn("w:numId"))] = abstract.get(aid, {})
            self.abstract_of[num.get(qn("w:numId"))] = aid

    def label(self, p) -> str:
        qn = self.qn
        num_pr = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if num_pr is None or num_pr.find(qn("w:numId")) is None:
            return ""
        num_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
        ilvl_el = num_pr.find(qn("w:ilvl"))
        ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
        levels = self.levels.get(num_id)
        if not levels or ilvl not in levels:
            return ""
        counters = self.counters.setdefault(self.abstract_of[num_id], [0] * 9)
        for i in range(ilvl + 1, 9):
            counters[i] = 0
        for i in range(ilvl):
            if counters[i] == 0 and i in levels:
                counters[i] = levels[i][2]
        counters[ilvl] = counters[ilvl] + 1 if counters[ilvl] else levels[ilvl][2]
        text = levels[ilvl][1]
        for i in range(ilvl + 1):
            fmt = levels.get(i, ("decimal", "", 1))[0]
            value = counters[i]
            rendered = LETTERS[(value - 1) % len(LETTERS)] if fmt.startswith("russian") else str(value)
            text = text.replace(f"%{i + 1}", rendered)
        return text + " " if text else ""


# --- clause parsing ---


def _is_successor(prev: tuple[int, ...] | None, cand: tuple[int, ...]) -> bool:
    if prev is None:
        return True
    for k in range(len(prev) + 1):
        if len(cand) <= k or cand[:k] != prev[:k]:
            continue
        base = prev[k] if k < len(prev) else 0
        if 1 <= cand[k] - base <= MAX_GAP and all(x == 1 for x in cand[k + 1:]):
            return True
    return False


def _line_at(text: str, pos: int) -> tuple[str, str]:
    """(text of the line before pos, text of the whole line containing pos)."""
    start = text.rfind("\n", 0, pos) + 1
    end = text.find("\n", pos)
    return text[start:pos], text[start:end if end != -1 else len(text)]


def _find_markers(text: str) -> list[tuple[int, int, str]]:
    """Accepted clause markers as (start, end, clause_id), in document order."""
    candidates = [(m.start(1), m.end(), "num", m.group(1)) for m in NUM_RE.finditer(text)]
    candidates += [(m.start(), m.end(), "sub", m.group(1)) for m in SUB_RE.finditer(text)]
    candidates += [(m.start(), m.end(), "toc", "toc") for m in TOC_RE.finditer(text)]
    candidates.sort()

    markers: list[tuple[int, int, str]] = []
    prev: tuple[int, ...] | None = None
    section: int | None = None
    next_letter = 0
    last_end = -1
    for start, end, kind, value in candidates:
        if start < last_end:
            continue
        before, line = _line_at(text, start)
        at_line_start = not before.strip()
        if kind == "toc":
            clause_id = "toc"
        elif kind == "sub":
            if prev is None or len(prev) < 2:
                continue
            idx = LETTERS.find(value)
            if idx not in (next_letter, next_letter + 1):
                continue
            next_letter = idx + 1
            clause_id = ".".join(map(str, prev)) + "." + value
        else:
            cand = tuple(int(x) for x in value.split("."))
            capital_next = bool(CAPITAL.match(text, end))
            if len(cand) == 1:
                if not capital_next or TOC_LINE.search(line):
                    continue
                if section is not None and cand[0] != section + 1:
                    continue
            elif not (at_line_start or capital_next) or not _is_successor(prev, cand):
                continue
            prev, section, next_letter = cand, cand[0], 0
            clause_id = value
        markers.append((start, end, clause_id))
        last_end = end
    return markers


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def parse_clauses(pages: list[str], paged: bool = True) -> list[Clause]:
    """Split page texts into clauses; page numbers are 1-based (None if not paged)."""
    text, offsets = "", []
    for page in pages:
        offsets.append(len(text))
        text += FOOTER.sub("", page.rstrip()) + "\n"

    def page_of(pos: int) -> int | None:
        if not paged:
            return None
        return max(i for i, off in enumerate(offsets) if off <= pos) + 1

    markers = _find_markers(text)
    clauses = []
    first = markers[0][0] if markers else len(text)
    if _clean(text[:first]):
        clauses.append(Clause(clause_id="preamble", text=_clean(text[:first]), page=page_of(0)))
    ids = set()
    for i, (start, end, clause_id) in enumerate(markers):
        stop = markers[i + 1][0] if i + 1 < len(markers) else len(text)
        clauses.append(
            Clause(clause_id=clause_id, text=_clean(text[end:stop]), page=page_of(start), parent_id=_parent(clause_id, ids))
        )
        ids.add(clause_id)
    if len(markers) < 3:  # no numbering: fall back to paragraphs so citations still work
        paragraphs = [p for p in re.split(r"\n\s*\n|\n(?=\s)", text) if _clean(p)]
        pos, clauses = 0, []
        for i, para in enumerate(paragraphs, 1):
            pos = text.find(para, pos)
            clauses.append(Clause(clause_id=f"p{i}", text=_clean(para), page=page_of(max(pos, 0))))
    return clauses


def _parent(clause_id: str, known: set[str]) -> str | None:
    parts = clause_id.split(".")
    for n in range(len(parts) - 1, 0, -1):
        candidate = ".".join(parts[:n])
        if candidate in known:
            return candidate
    return None


# --- entry points ---


def ingest_file(path: str | Path, side: Side, doc_id: str, name: str | None = None) -> Document:
    path = Path(path)
    name = name or path.name
    suffix = path.suffix.lower()
    warnings: list[str] = []
    try:
        if suffix == ".pdf":
            pages = pdf_pages(path)
            for i, page in enumerate(pages, 1):
                if len(page.strip()) < 20:
                    warnings.append(f"Страница {i}: нет текстового слоя (скан?), текст не извлечён")
            clauses, n_pages = parse_clauses(pages), len(pages)
        elif suffix == ".docx":
            clauses, n_pages = parse_clauses([docx_text(path)], paged=False), None
        elif suffix in (".txt", ".md"):
            clauses, n_pages = parse_clauses([path.read_text(encoding="utf-8")], paged=False), None
        else:
            raise ValueError(f"неподдерживаемый формат {suffix}")
    except Exception as exc:  # a broken file must not kill the whole run
        return Document(doc_id=doc_id, name=name, side=side, pages=None, clauses=[],
                        warnings=[f"Не удалось прочитать файл: {exc}"])
    if not clauses:
        warnings.append("В документе не найден текст")
    return Document(doc_id=doc_id, name=name, side=side, pages=n_pages, clauses=clauses, warnings=warnings)


def ingest_set(paths: list[str | Path], side: Side) -> list[Document]:
    return [ingest_file(p, side, doc_id=f"{side}-{i}") for i, p in enumerate(paths, 1)]


def dump_text(path: str | Path) -> str:
    """Page-marked plain text, for grepping (`data/samples/*.txt`)."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        return "".join(f"\n=== page {i} ===\n{t}" for i, t in enumerate(pdf_pages(path), 1)).lstrip()
    return docx_text(path) if path.suffix.lower() == ".docx" else path.read_text(encoding="utf-8")


if __name__ == "__main__":
    # python -m backend.ingest file.pdf [...]  -> writes file.txt next to each, prints clause stats
    for arg in sys.argv[1:]:
        src = Path(arg)
        src.with_suffix(".txt").write_text(dump_text(src), encoding="utf-8")
        doc = ingest_file(src, "before", doc_id=src.stem)
        print(f"{src}: {len(doc.clauses)} clauses, {doc.pages} pages, warnings={doc.warnings}")
